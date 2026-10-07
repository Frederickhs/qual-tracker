from pathlib import Path
import re
import sys
import time
import warnings

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
QUALIFICATION_FOLDER = PROJECT_ROOT / "Input" / "Qualification"
TASK_FOLDER = PROJECT_ROOT / "Input" / "Tasks"
CONFIG_FILE = PROJECT_ROOT / "Config" / "qualification_mapping.csv"
LOG_FOLDER = PROJECT_ROOT / "Logs"
VALIDATION_FILE = LOG_FOLDER / "cycle_validation.csv"
SUMMARY_FILE = LOG_FOLDER / "cycle_summary.csv"

TODAY = pd.Timestamp.today().normalize()
DUE_SOON_DAYS = 30
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}

REQUIRED_TASK_COLUMNS = {
    "Task Num",
    "Status",
    "Changed Date",
    "Change By",
    "Name",
    "Qualification",
}

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style*",
    category=UserWarning,
)


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def excel_files(folder):
    return sorted(
        [
            path
            for path in folder.iterdir()
            if path.is_file()
            and not path.name.startswith("~$")
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ],
        key=lambda path: path.name.lower(),
    )


def load_mapping():
    mapping = pd.read_csv(CONFIG_FILE, dtype=str).fillna("")
    mapping.columns = [normalize_text(column) for column in mapping.columns]

    required = {"Area", "Venue", "Qualification", "Discipline", "Tracked"}
    missing = required.difference(mapping.columns)
    if missing:
        raise ValueError(
            "Configuration is missing columns: " + ", ".join(sorted(missing))
        )

    for column in required:
        mapping[column] = mapping[column].astype(str).str.strip()

    mapping["Qualification"] = mapping["Qualification"].str.upper()
    mapping["Tracked"] = mapping["Tracked"].str.lower()
    return mapping[mapping["Tracked"] == "yes"].copy()


def qualification_from_filename(file_path, approved_codes):
    filename = file_path.stem.upper()
    matches = [code for code in approved_codes if code in filename]

    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValueError("qualification code was not found in filename")
    raise ValueError("more than one qualification code was found in filename")


def read_qualifications(mapping):
    files = excel_files(QUALIFICATION_FOLDER)
    approved_codes = set(mapping["Qualification"])
    reports = []
    detected_codes = []

    if not files:
        raise ValueError("No qualification reports were found")

    for file_path in files:
        qualification = qualification_from_filename(file_path, approved_codes)
        detected_codes.append(qualification)

        mapping_rows = mapping[mapping["Qualification"] == qualification]
        if len(mapping_rows) != 1:
            raise ValueError(
                f"{qualification} must have exactly one configuration row"
            )
        mapping_row = mapping_rows.iloc[0]

        workbook = pd.ExcelFile(file_path, engine="openpyxl")
        if "List of Labor" in workbook.sheet_names:
            sheet_name = "List of Labor"
        elif len(workbook.sheet_names) == 1:
            sheet_name = workbook.sheet_names[0]
        else:
            raise ValueError(
                f"{file_path.name}: worksheet 'List of Labor' was not found"
            )

        report = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            engine="openpyxl",
            dtype={"Labor": "string"},
        )
        report.columns = [normalize_text(column) for column in report.columns]

        required = {"Labor", "Name", "Effective Date", "Status"}
        missing = required.difference(report.columns)
        if missing:
            raise ValueError(
                f"{file_path.name}: missing columns: {', '.join(sorted(missing))}"
            )

        report = report.dropna(how="all").copy()
        report["SSO"] = report["Labor"].astype("string").str.strip()
        report["Technician Name"] = report["Name"].astype("string").str.strip()
        report["Qualification Date"] = pd.to_datetime(
            report["Effective Date"], format="mixed", errors="coerce"
        ).dt.normalize()
        report["Qualification Status"] = (
            report["Status"].astype("string").str.strip().str.upper()
        )
        report["Qualification"] = qualification
        report["Area"] = mapping_row["Area"]
        report["Venue"] = mapping_row["Venue"]
        report["Discipline"] = mapping_row["Discipline"]
        report["Qualification Source File"] = file_path.name

        reports.append(
            report[
                [
                    "Area",
                    "Venue",
                    "Discipline",
                    "Qualification",
                    "SSO",
                    "Technician Name",
                    "Qualification Date",
                    "Qualification Status",
                    "Qualification Source File",
                ]
            ]
        )

    duplicate_codes = sorted(
        {code for code in detected_codes if detected_codes.count(code) > 1}
    )
    if duplicate_codes:
        raise ValueError(
            "Duplicate qualification reports: " + ", ".join(duplicate_codes)
        )

    missing_codes = sorted(approved_codes.difference(detected_codes))
    if missing_codes:
        raise ValueError(
            "Missing qualification reports: " + ", ".join(missing_codes)
        )

    combined = pd.concat(reports, ignore_index=True)

    if combined["Qualification Date"].isna().any():
        count = int(combined["Qualification Date"].isna().sum())
        raise ValueError(f"Invalid qualification dates found: {count}")

    duplicate_mask = combined.duplicated(
        subset=["SSO", "Qualification"], keep=False
    )
    if duplicate_mask.any():
        count = int(duplicate_mask.sum())
        raise ValueError(f"Duplicate SSO and qualification rows found: {count}")

    return combined


def find_header_row(worksheet, maximum_rows=15):
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=maximum_rows,
            values_only=True,
        ),
        start=1,
    ):
        headers = {
            normalize_text(value): index
            for index, value in enumerate(row)
            if normalize_text(value)
        }
        if REQUIRED_TASK_COLUMNS.issubset(headers):
            return row_number, headers

    return None, None


def read_tasks():
    files = excel_files(TASK_FOLDER)
    rows = []

    if not files:
        raise ValueError("No task reports were found")

    for file_number, file_path in enumerate(files, start=1):
        started = time.perf_counter()
        print(f"[{file_number}/{len(files)}] Reading {file_path.name}...")

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        file_count = 0

        for worksheet in workbook.worksheets:
            worksheet.reset_dimensions()
            header_row_number, headers = find_header_row(worksheet)
            if headers is None:
                continue

            highest_index = max(headers.values())

            for row in worksheet.iter_rows(
                min_row=header_row_number + 1,
                values_only=True,
            ):
                if len(row) <= highest_index:
                    continue

                status = normalize_text(row[headers["Status"]]).upper()
                if status != "COMP":
                    continue

                qualification = normalize_text(
                    row[headers["Qualification"]]
                ).upper()
                sso = normalize_text(row[headers["Change By"]])
                completion_date = row[headers["Changed Date"]]
                task_number = normalize_text(row[headers["Task Num"]])
                technician_name = normalize_text(row[headers["Name"]])
                task_description = normalize_text(row[0]) if row else ""

                rows.append(
                    {
                        "SSO": sso,
                        "Task Technician Name": technician_name,
                        "Qualification": qualification,
                        "PM Completion Date": completion_date,
                        "Task Number": task_number,
                        "Task Description": task_description,
                        "Task Source File": file_path.name,
                        "Task Source Sheet": worksheet.title,
                    }
                )
                file_count += 1

        workbook.close()
        seconds = time.perf_counter() - started
        print(f"  COMP rows collected: {file_count:,}")
        print(f"  Processing time: {seconds:,.1f} seconds")

    tasks = pd.DataFrame(rows)
    if tasks.empty:
        return tasks, set()

    tasks["SSO"] = tasks["SSO"].astype(str).str.strip()
    tasks["Qualification"] = (
        tasks["Qualification"].astype(str).str.strip().str.upper()
    )
    tasks["PM Completion Date"] = pd.to_datetime(
        tasks["PM Completion Date"], format="mixed", errors="coerce"
    )

    tasks = tasks.dropna(subset=["PM Completion Date"])
    tasks = tasks[tasks["SSO"] != ""].copy()
    tasks = tasks.drop_duplicates(
        subset=[
            "SSO",
            "Qualification",
            "PM Completion Date",
            "Task Number",
            "Task Description",
        ]
    )

    qualifications_found = set(tasks["Qualification"].dropna().unique())
    return tasks, qualifications_found


def add_years(date_value, years):
    return date_value + pd.DateOffset(years=int(years))


def build_cycles(qualifications):
    cycle_rows = []

    for record in qualifications.to_dict("records"):
        start = pd.Timestamp(record["Qualification Date"]).normalize()
        cycle_number = 1

        while start <= TODAY:
            due = add_years(start, 1)
            cycle_record = record.copy()
            cycle_record["Cycle Number"] = cycle_number
            cycle_record["Cycle Start Date"] = start
            cycle_record["Cycle Due Date"] = due
            cycle_rows.append(cycle_record)

            cycle_number += 1
            start = add_years(record["Qualification Date"], cycle_number - 1)

    return pd.DataFrame(cycle_rows)


def assign_tasks_to_cycles(cycles, tasks, qualifications_with_tasks):
    task_groups = {
        key: group.sort_values("PM Completion Date")
        for key, group in tasks.groupby(["SSO", "Qualification"])
    }

    output_rows = []

    for cycle in cycles.to_dict("records"):
        key = (str(cycle["SSO"]), cycle["Qualification"])
        task_data_available = cycle["Qualification"] in qualifications_with_tasks
        latest_task = None

        if task_data_available and key in task_groups:
            candidates = task_groups[key]
            valid = candidates[
                (candidates["PM Completion Date"].dt.normalize() >= cycle["Cycle Start Date"])
                & (candidates["PM Completion Date"].dt.normalize() <= cycle["Cycle Due Date"])
            ]
            if not valid.empty:
                latest_task = valid.iloc[-1]

        days_remaining = int((cycle["Cycle Due Date"] - TODAY).days)

        if not task_data_available:
            result = "Task Data Missing"
        elif latest_task is not None:
            result = "Good / Completed"
        elif days_remaining < 0:
            result = "Requalification Needed"
        elif days_remaining <= DUE_SOON_DAYS:
            result = "Due Soon"
        else:
            result = "Good"

        output = cycle.copy()
        output["Task Data Available"] = "Yes" if task_data_available else "No"
        output["PM Completion Date"] = (
            latest_task["PM Completion Date"] if latest_task is not None else pd.NaT
        )
        output["Task Description"] = (
            latest_task["Task Description"] if latest_task is not None else ""
        )
        output["Task Number"] = (
            latest_task["Task Number"] if latest_task is not None else ""
        )
        output["Task Source File"] = (
            latest_task["Task Source File"] if latest_task is not None else ""
        )
        output["Days Remaining"] = days_remaining
        output["Cycle Result"] = result
        output_rows.append(output)

    result = pd.DataFrame(output_rows)

    columns = [
        "Area",
        "Venue",
        "Discipline",
        "Qualification",
        "SSO",
        "Technician Name",
        "Qualification Date",
        "Qualification Status",
        "Cycle Number",
        "Cycle Start Date",
        "Cycle Due Date",
        "Task Data Available",
        "PM Completion Date",
        "Task Description",
        "Task Number",
        "Days Remaining",
        "Cycle Result",
        "Qualification Source File",
        "Task Source File",
    ]

    return result[columns].sort_values(
        ["Venue", "Qualification", "Technician Name", "Cycle Number"],
        kind="stable",
    )


def main():
    started = time.perf_counter()

    print("=" * 70)
    print("PM Qualification Activity Tracker - Cycle Processor")
    print("=" * 70)
    print(f"Processing date: {TODAY.date()}")

    try:
        mapping = load_mapping()
        qualifications = read_qualifications(mapping)
    except Exception as error:
        print(f"ERROR reading qualification reports: {error}")
        return 1

    print(f"Qualification records loaded: {len(qualifications):,}")

    try:
        tasks, qualifications_with_tasks = read_tasks()
    except Exception as error:
        print(f"ERROR reading task reports: {error}")
        return 1

    print(f"Completed task rows loaded: {len(tasks):,}")

    cycles = build_cycles(qualifications)
    results = assign_tasks_to_cycles(
        cycles,
        tasks,
        qualifications_with_tasks,
    )

    LOG_FOLDER.mkdir(parents=True, exist_ok=True)
    results.to_csv(VALIDATION_FILE, index=False, date_format="%Y-%m-%d %H:%M:%S")

    summary = (
        results.groupby(
            ["Venue", "Qualification", "Discipline", "Cycle Result"],
            dropna=False,
        )
        .size()
        .reset_index(name="Cycle Records")
        .sort_values(["Venue", "Qualification", "Cycle Result"])
    )
    summary.to_csv(SUMMARY_FILE, index=False)

    print()
    print("Summary")
    print("-------")
    print(f"Qualification records: {len(qualifications):,}")
    print(f"Annual cycle records: {len(results):,}")
    print(f"Completed task rows available: {len(tasks):,}")
    print(f"Technicians: {results['SSO'].nunique():,}")

    print()
    print("Cycle results")
    print("-------------")
    for status, count in results["Cycle Result"].value_counts().items():
        print(f"{status}: {count:,}")

    missing = sorted(
        set(mapping["Qualification"]).difference(qualifications_with_tasks)
    )
    if missing:
        print()
        print("Qualifications with missing task data")
        print("-------------------------------------")
        for qualification in missing:
            print(f"- {qualification}")

    print()
    print(f"Validation file: {VALIDATION_FILE}")
    print(f"Summary file: {SUMMARY_FILE}")
    print(f"Total processing time: {time.perf_counter() - started:,.1f} seconds")
    print()
    print("Cycle processing completed successfully.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
