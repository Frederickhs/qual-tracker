from pathlib import Path
import sys
import time
import warnings

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
TASK_FOLDER = PROJECT_ROOT / "Input" / "Tasks"
CONFIG_FILE = PROJECT_ROOT / "Config" / "qualification_mapping.csv"
LOG_FOLDER = PROJECT_ROOT / "Logs"
SUMMARY_FILE = LOG_FOLDER / "task_summary.csv"

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
REQUIRED_TASK_COLUMNS = {
    "Task Num",
    "Status",
    "Changed Date",
    "Change By",
    "Name",
    "Qualification",
}

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style*",
    category=UserWarning,
)


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def get_task_files():
    return sorted(
        [
            file_path
            for file_path in TASK_FOLDER.iterdir()
            if file_path.is_file()
            and not file_path.name.startswith("~$")
            and file_path.suffix.lower() in SUPPORTED_EXTENSIONS
        ],
        key=lambda file_path: file_path.name.lower(),
    )


def load_tracked_qualifications():
    mapping = pd.read_csv(CONFIG_FILE, dtype=str).fillna("")
    mapping.columns = [normalize_text(column) for column in mapping.columns]

    required = {"Qualification", "Tracked"}
    missing = required.difference(mapping.columns)

    if missing:
        raise ValueError(
            "Configuration is missing columns: " + ", ".join(sorted(missing))
        )

    mapping["Qualification"] = (
        mapping["Qualification"].astype(str).str.strip().str.upper()
    )
    mapping["Tracked"] = mapping["Tracked"].astype(str).str.strip().str.lower()

    return set(
        mapping.loc[mapping["Tracked"] == "yes", "Qualification"]
    )


def find_header_row(worksheet, maximum_rows=15):
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=maximum_rows,
            values_only=True,
        ),
        start=1,
    ):
        headers = {
            normalize_text(value): index
            for index, value in enumerate(row)
            if normalize_text(value)
        }

        if REQUIRED_TASK_COLUMNS.issubset(headers):
            return row_number, headers

    return None, None


def read_completed_tasks(file_path):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    completed_rows = []
    sheets_with_data = 0

    for worksheet in workbook.worksheets:
        worksheet.reset_dimensions()
        header_row_number, headers = find_header_row(worksheet)

        if headers is None:
            continue

        sheets_with_data += 1
        highest_index = max(headers.values())

        for row in worksheet.iter_rows(
            min_row=header_row_number + 1,
            values_only=True,
        ):
            if len(row) <= highest_index:
                continue

            status = normalize_text(row[headers["Status"]]).upper()

            if status != "COMP":
                continue

            qualification = normalize_text(
                row[headers["Qualification"]]
            ).upper()
            sso = normalize_text(row[headers["Change By"]])
            completion_date = row[headers["Changed Date"]]
            task_number = normalize_text(row[headers["Task Num"]])
            technician_name = normalize_text(row[headers["Name"]])

            first_value = normalize_text(row[0]) if row else ""
            task_description = first_value

            completed_rows.append(
                {
                    "SSO": sso,
                    "Technician Name": technician_name,
                    "Qualification": qualification,
                    "PM Completion Date": completion_date,
                    "Task Number": task_number,
                    "Task Description": task_description,
                    "Source File": file_path.name,
                    "Source Sheet": worksheet.title,
                }
            )

    workbook.close()
    return completed_rows, sheets_with_data


def main():
    started = time.perf_counter()

    print("=" * 70)
    print("PM Qualification Activity Tracker - Completed Task Reader")
    print("=" * 70)

    if not CONFIG_FILE.exists():
        print(f"ERROR: Configuration file not found: {CONFIG_FILE}")
        return 1

    task_files = get_task_files()
    print(f"Task reports found: {len(task_files)}")

    if not task_files:
        print("ERROR: No task reports were found.")
        return 1

    try:
        tracked_qualifications = load_tracked_qualifications()
    except Exception as error:
        print(f"ERROR: {error}")
        return 1

    all_completed_rows = []
    file_results = []

    for file_number, file_path in enumerate(task_files, start=1):
        print()
        print(f"[{file_number}/{len(task_files)}] Reading {file_path.name}...")
        file_started = time.perf_counter()

        try:
            rows, sheets_with_data = read_completed_tasks(file_path)
        except Exception as error:
            print(f"ERROR: {file_path.name}: {error}")
            return 1

        file_seconds = time.perf_counter() - file_started
        qualifications = sorted(
            {
                row["Qualification"]
                for row in rows
                if row["Qualification"]
            }
        )

        all_completed_rows.extend(rows)
        file_results.append(
            {
                "Source File": file_path.name,
                "Worksheets With Data": sheets_with_data,
                "COMP Rows": len(rows),
                "Qualifications": ", ".join(qualifications),
                "Processing Seconds": round(file_seconds, 1),
            }
        )

        print(f"  Worksheets processed: {sheets_with_data:,}")
        print(f"  COMP rows collected: {len(rows):,}")
        print(
            "  Qualifications: "
            + (", ".join(qualifications) if qualifications else "None")
        )
        print(f"  Processing time: {file_seconds:,.1f} seconds")

    completed = pd.DataFrame(all_completed_rows)

    if completed.empty:
        print()
        print("ERROR: No completed task rows were found.")
        return 1

    completed["SSO"] = completed["SSO"].astype(str).str.strip()
    completed["Qualification"] = (
        completed["Qualification"].astype(str).str.strip().str.upper()
    )
    completed["PM Completion Date"] = pd.to_datetime(
        completed["PM Completion Date"],
        format="mixed",
        errors="coerce",
    )

    exact_duplicate_mask = completed.duplicated(
        subset=[
            "SSO",
            "Qualification",
            "PM Completion Date",
            "Task Number",
            "Task Description",
        ],
        keep=False,
    )

    qualifications_found = set(
        completed.loc[
            completed["Qualification"] != "",
            "Qualification",
        ]
    )
    missing_qualifications = sorted(
        tracked_qualifications.difference(qualifications_found)
    )
    unknown_qualifications = sorted(
        qualifications_found.difference(tracked_qualifications)
    )

    summary = (
        completed.groupby("Qualification", dropna=False)
        .agg(
            COMP_Rows=("Task Number", "size"),
            Unique_Technicians=("SSO", "nunique"),
            Earliest_Completion=("PM Completion Date", "min"),
            Latest_Completion=("PM Completion Date", "max"),
            Invalid_Dates=(
                "PM Completion Date",
                lambda values: int(values.isna().sum()),
            ),
        )
        .reset_index()
        .sort_values("Qualification")
    )

    LOG_FOLDER.mkdir(parents=True, exist_ok=True)
    summary.to_csv(SUMMARY_FILE, index=False)

    total_seconds = time.perf_counter() - started

    print()
    print("Summary")
    print("-------")
    print(f"Task reports processed: {len(task_files)}")
    print(f"COMP rows combined: {len(completed):,}")
    print(f"Unique task technicians: {completed['SSO'].nunique():,}")
    print(
        "Exact duplicate task rows: "
        f"{int(exact_duplicate_mask.sum()):,}"
    )
    print(
        "Invalid completion dates: "
        f"{int(completed['PM Completion Date'].isna().sum()):,}"
    )
    print(f"Total processing time: {total_seconds:,.1f} seconds")
    print(f"Summary file created: {SUMMARY_FILE}")

    print()
    print("Qualifications found")
    print("--------------------")

    for qualification in sorted(qualifications_found):
        print(f"- {qualification}")

    if missing_qualifications:
        print()
        print("WARNING: Missing tracked task qualifications")
        print("--------------------------------------------")
        for qualification in missing_qualifications:
            print(f"- {qualification}")

    if unknown_qualifications:
        print()
        print("ERROR: Unknown task qualifications")
        print("----------------------------------")
        for qualification in unknown_qualifications:
            print(f"- {qualification}")
        return 1

    print()
    print("Completed-task reading finished successfully.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
