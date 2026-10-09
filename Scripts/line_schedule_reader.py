from pathlib import Path
import json
import sys

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
SETTINGS_FILE = PROJECT_ROOT / "Config" / "settings.json"

SHEET_SHIFT_MAPPING = {
    "1st Shift": "1st",
    "2nd Shift": "2nd",
    "3rd Shift Unchained": "3rd",
    "3rd Shift Wolf": "3rd",
}


def normalize_sso(value):
    if value is None:
        return ""

    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    return text


def load_settings():
    if not SETTINGS_FILE.exists():
        raise FileNotFoundError(
            f"Settings file not found: {SETTINGS_FILE}"
        )

    with SETTINGS_FILE.open("r", encoding="utf-8-sig") as file:
        settings = json.load(file)

    folder_value = settings.get("line_schedule_folder", "").strip()

    if not folder_value:
        raise ValueError(
            "settings.json does not contain line_schedule_folder."
        )

    return Path(folder_value)


def find_schedule_workbook(folder):
    if not folder.exists():
        raise FileNotFoundError(
            f"Line Schedule folder not found: {folder}"
        )

    files = sorted(
        [
            file_path
            for file_path in folder.glob("*.xlsx")
            if not file_path.name.startswith("~$")
        ],
        key=lambda file_path: file_path.name.lower(),
    )

    if not files:
        raise FileNotFoundError(
            f"No Line Schedule workbook found in: {folder}"
        )

    if len(files) > 1:
        names = ", ".join(file_path.name for file_path in files)

        raise ValueError(
            "More than one Line Schedule workbook was found: "
            f"{names}"
        )

    return files[0]


def read_line_schedule():
    schedule_folder = load_settings()
    workbook_path = find_schedule_workbook(schedule_folder)

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    records = []

    try:
        for worksheet in workbook.worksheets:
            worksheet.reset_dimensions()

            shift = SHEET_SHIFT_MAPPING.get(worksheet.title)

            if shift is None:
                continue

            for row in worksheet.iter_rows(
                min_row=4,
                values_only=True,
            ):
                technician_name = (
                    str(row[0]).strip()
                    if len(row) > 0 and row[0] is not None
                    else ""
                )

                sso = (
                    normalize_sso(row[3])
                    if len(row) > 3
                    else ""
                )

                if not sso:
                    continue

                records.append(
                    {
                        "SSO": sso,
                        "Schedule Technician Name": technician_name,
                        "Shift": shift,
                        "Schedule Sheet": worksheet.title,
                        "Schedule Source File": workbook_path.name,
                    }
                )

    finally:
        workbook.close()

    schedule = pd.DataFrame(records)

    if schedule.empty:
        raise ValueError(
            "No technicians with SSOs were found in the Line Schedule."
        )

    duplicate_mask = schedule.duplicated(
        subset=["SSO"],
        keep=False,
    )

    duplicate_rows = schedule[duplicate_mask].copy()
    conflicting_ssos = set()

    if not duplicate_rows.empty:
        shift_counts = duplicate_rows.groupby("SSO")["Shift"].nunique()

        conflicting_ssos = set(
            shift_counts[shift_counts > 1].index
        )

    schedule["Shift Conflict"] = schedule["SSO"].apply(
        lambda sso: "Yes" if sso in conflicting_ssos else "No"
    )

    schedule = schedule.sort_values(
        ["Shift", "Schedule Sheet", "Schedule Technician Name"],
        kind="stable",
    )

    return schedule, workbook_path, conflicting_ssos


def main():
    print("=" * 70)
    print("PM Qualification Activity Tracker - Line Schedule Reader")
    print("=" * 70)

    try:
        schedule, workbook_path, conflicting_ssos = (
            read_line_schedule()
        )
    except Exception as error:
        print(f"ERROR: {error}")
        return 1

    print(f"Line Schedule: {workbook_path}")
    print(f"Schedule technicians: {schedule['SSO'].nunique():,}")

    print()
    print("Shift counts")
    print("------------")

    for shift, count in schedule.groupby("Shift")["SSO"].nunique().items():
        print(f"{shift}: {count:,}")

    print()
    print(
        "Duplicate schedule rows: "
        f"{schedule.duplicated('SSO', keep=False).sum():,}"
    )
    print(f"Conflicting shift SSOs: {len(conflicting_ssos):,}")

    if conflicting_ssos:
        print()
        print("SSOs with conflicting shifts")
        print("----------------------------")

        for sso in sorted(conflicting_ssos):
            shifts = sorted(
                schedule.loc[schedule["SSO"] == sso, "Shift"].unique()
            )

            print(f"{sso}: {', '.join(shifts)}")

    print()
    print("Line Schedule reading completed successfully.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
