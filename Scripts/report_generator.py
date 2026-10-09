from pathlib import Path
import re
import sys
import time
from cycle_history import apply_cycle_history, save_cycle_history
from line_schedule_reader import read_line_schedule

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

from cycle_processor import (
    TODAY,
    assign_tasks_to_cycles,
    build_cycles,
    load_mapping,
    read_qualifications,
    read_tasks,
)


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FOLDER = PROJECT_ROOT / "Output"
MANUAL_ROSTER_FOLDER = PROJECT_ROOT / "Input" / "Manual roster"
ANIMATION_ROSTER_FILE = MANUAL_ROSTER_FOLDER / "Animation Roster.xlsx"

OUTPUT_FILES = {
    "Wolf": "Wolf Qualification Tracker.xlsx",
    "Darkmoor": "Darkmoor Qualification Tracker.xlsx",
    "Unchained": "Unchained Qualification Tracker.xlsx",
    "Animation": "Animation Qualification Tracker.xlsx",
}

STATUS_STYLES = {
    "Good / Completed": ("C6EFCE", "006100"),
    "Good": ("DDEBF7", "1F4E78"),
    "Due Soon": ("FFF2CC", "7F6000"),
    "Requalification Needed": ("F4CCCC", "9C0006"),
    "Task Data Missing": ("E4DFEC", "5F497A"),
}

DISPLAY_COLUMNS = [
    "Venue",
    "Shift",
    "Discipline",
    "Qualification",
    "SSO",
    "Technician Name",
    "Qualification Date",
    "Qualification Status",
    "Cycle Number",
    "Cycle Start Date",
    "Cycle Due Date",
    "Task Data Available",
    "PM Completion Date",
    "Task Description",
    "Task Number",
    "Days Remaining",
    "Cycle Result",
]

ROSTER_EXCEPTION_COLUMNS = [
    "Exception Type",
    "Roster Name",
    "SSO",
    "Official Maximo Name",
    "Details",
]


def normalize_sso(value):
    if pd.isna(value):
        return ""

    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def normalized_name_tokens(value):
    if pd.isna(value):
        return []

    text = str(value).upper().replace("(UDX)", " ")
    return sorted(re.findall(r"[A-Z0-9]+", text))


def read_animation_roster(qualifications):
    if not ANIMATION_ROSTER_FILE.exists():
        raise ValueError(
            "Animation roster was not found. Expected: "
            f"{ANIMATION_ROSTER_FILE}"
        )

    roster = pd.read_excel(
        ANIMATION_ROSTER_FILE,
        engine="openpyxl",
        dtype={"ID": "string", "Name": "string"},
    )
    roster.columns = [str(column).strip() for column in roster.columns]

    required_columns = {"Name", "ID"}
    missing_columns = required_columns.difference(roster.columns)
    if missing_columns:
        raise ValueError(
            "Animation roster is missing columns: "
            + ", ".join(sorted(missing_columns))
        )

    roster = roster.dropna(how="all").copy()
    roster["Roster Name"] = roster["Name"].fillna("").astype(str).str.strip()
    roster["SSO"] = roster["ID"].apply(normalize_sso)

    official = qualifications[
        qualifications["Qualification"] == "UCMOD4"
    ].copy()
    official_by_sso = {
        str(row["SSO"]): row
        for row in official.to_dict("records")
    }

    exceptions = []

    for row_number, row in roster.iterrows():
        if not row["SSO"]:
            exceptions.append(
                {
                    "Exception Type": "Blank or invalid SSO",
                    "Roster Name": row["Roster Name"],
                    "SSO": "",
                    "Official Maximo Name": "",
                    "Details": f"Roster row {row_number + 2} has no usable ID.",
                }
            )

    nonblank = roster[roster["SSO"] != ""].copy()
    duplicate_mask = nonblank.duplicated(subset=["SSO"], keep=False)

    for _, row in nonblank[duplicate_mask].iterrows():
        exceptions.append(
            {
                "Exception Type": "Duplicate roster SSO",
                "Roster Name": row["Roster Name"],
                "SSO": row["SSO"],
                "Official Maximo Name": "",
                "Details": "The SSO appears more than once in Animation Roster.xlsx.",
            }
        )

    unique_roster = nonblank.drop_duplicates(subset=["SSO"], keep="first")
    matched_ssos = set()

    for _, row in unique_roster.iterrows():
        sso = row["SSO"]
        official_row = official_by_sso.get(sso)

        if official_row is None:
            exceptions.append(
                {
                    "Exception Type": "SSO not found in UCMOD4",
                    "Roster Name": row["Roster Name"],
                    "SSO": sso,
                    "Official Maximo Name": "",
                    "Details": (
                        "The roster SSO was not found in the official "
                        "UCMOD4 qualification report."
                    ),
                }
            )
            continue

        matched_ssos.add(sso)
        official_name = str(official_row["Technician Name"])

        if normalized_name_tokens(row["Roster Name"]) != normalized_name_tokens(
            official_name
        ):
            exceptions.append(
                {
                    "Exception Type": "Roster name differs from Maximo",
                    "Roster Name": row["Roster Name"],
                    "SSO": sso,
                    "Official Maximo Name": official_name,
                    "Details": (
                        "The official Maximo name is used in the generated report."
                    ),
                }
            )

    exception_frame = pd.DataFrame(exceptions, columns=ROSTER_EXCEPTION_COLUMNS)
    return matched_ssos, exception_frame, len(unique_roster)


def add_shift_to_results(results):
    schedule, workbook_path, conflicting_ssos = read_line_schedule()

    usable_schedule = schedule[
        schedule["Shift Conflict"] != "Yes"
    ].drop_duplicates(subset=["SSO"], keep="first")

    shift_by_sso = usable_schedule.set_index("SSO")["Shift"].to_dict()

    updated = results.copy()
    updated["SSO"] = updated["SSO"].astype(str).str.strip()
    updated["Shift"] = updated["SSO"].map(shift_by_sso)
    updated["Shift"] = updated["Shift"].fillna("Shift Not Found")

    if conflicting_ssos:
        updated.loc[
            updated["SSO"].isin(conflicting_ssos),
            "Shift",
        ] = "Shift Conflict"

    stats = {
        "Line Schedule technicians": int(schedule["SSO"].nunique()),
        "Tracker technicians matched to a shift": int(
            updated.loc[
                updated["Shift"].isin(["1st", "2nd", "3rd"]),
                "SSO",
            ].nunique()
        ),
        "Tracker technicians with shift not found": int(
            updated.loc[
                updated["Shift"] == "Shift Not Found",
                "SSO",
            ].nunique()
        ),
        "Tracker technicians with shift conflict": int(
            updated.loc[
                updated["Shift"] == "Shift Conflict",
                "SSO",
            ].nunique()
        ),
    }

    return updated, workbook_path, stats


def write_dataframe(worksheet, dataframe):
    worksheet.append(list(dataframe.columns))

    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append(list(row))


def format_table_sheet(worksheet, status_column="Q"):
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="548235")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 15,
        "B": 14,
        "C": 15,
        "D": 15,
        "E": 14,
        "F": 31,
        "G": 18,
        "H": 20,
        "I": 14,
        "J": 18,
        "K": 18,
        "L": 19,
        "M": 21,
        "N": 44,
        "O": 16,
        "P": 16,
        "Q": 26,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }

    for row in worksheet.iter_rows(min_row=2):
        if "SSO" in headers:
            row[headers["SSO"] - 1].number_format = "@"

        for date_header in [
            "Qualification Date",
            "Cycle Start Date",
            "Cycle Due Date",
        ]:
            if date_header in headers:
                row[headers[date_header] - 1].number_format = "m/d/yyyy"

        if "PM Completion Date" in headers:
            row[headers["PM Completion Date"] - 1].number_format = (
                "m/d/yyyy h:mm AM/PM"
            )

        if "Task Description" in headers:
            row[headers["Task Description"] - 1].alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    if worksheet.max_row >= 2:
        worksheet.auto_filter.ref = worksheet.dimensions

        status_range = f"{status_column}2:{status_column}{worksheet.max_row}"
        for status, (fill_color, font_color) in STATUS_STYLES.items():
            worksheet.conditional_formatting.add(
                status_range,
                FormulaRule(
                    formula=[f'${status_column}2="{status}"'],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )


def create_dashboard(workbook, tracker_data, report_name):
    worksheet = workbook.active
    worksheet.title = "Dashboard"
    worksheet.sheet_view.showGridLines = False

    worksheet["A1"] = "PM Qualification Activity Tracker"
    worksheet["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.merge_cells("A1:D1")

    worksheet["A2"] = f"{report_name} Annual PM Compliance"
    worksheet["A2"].font = Font(size=12, bold=True)
    worksheet.merge_cells("A2:D2")

    worksheet["A4"] = "Metric"
    worksheet["B4"] = "Count"

    metrics = [
        ("Total Cycles", len(tracker_data)),
        (
            "Good / Completed",
            int((tracker_data["Cycle Result"] == "Good / Completed").sum()),
        ),
        ("Good", int((tracker_data["Cycle Result"] == "Good").sum())),
        ("Due Soon", int((tracker_data["Cycle Result"] == "Due Soon").sum())),
        (
            "Requalification Needed",
            int(
                (
                    tracker_data["Cycle Result"]
                    == "Requalification Needed"
                ).sum()
            ),
        ),
        (
            "Task Data Missing",
            int((tracker_data["Cycle Result"] == "Task Data Missing").sum()),
        ),
    ]

    for row_number, (label, value) in enumerate(metrics, start=5):
        worksheet.cell(row=row_number, column=1, value=label)
        worksheet.cell(row=row_number, column=2, value=value)
        worksheet.cell(row=row_number, column=1).font = Font(bold=True)
        worksheet.cell(row=row_number, column=2).font = Font(bold=True, size=14)

        if label in STATUS_STYLES:
            fill_color, font_color = STATUS_STYLES[label]
            worksheet.cell(row=row_number, column=1).fill = PatternFill(
                "solid", fgColor=fill_color
            )
            worksheet.cell(row=row_number, column=1).font = Font(
                bold=True, color=font_color
            )

    worksheet["A12"] = "Processing Date"
    worksheet["B12"] = TODAY.to_pydatetime()
    worksheet["B12"].number_format = "m/d/yyyy"

    worksheet["A14"] = (
        "Each PM counts only for the annual cycle in which it was completed."
    )
    worksheet["A14"].font = Font(italic=True, color="666666")
    worksheet.merge_cells("A14:F14")

    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 18


def create_exceptions_sheet(workbook, tracker_data):
    worksheet = workbook.create_sheet("Exceptions")

    exceptions = tracker_data[
        (tracker_data["Cycle Result"] == "Task Data Missing")
        | (tracker_data["Qualification Status"] != "ACTIVE")
        | (tracker_data["Shift"].isin(["Shift Not Found", "Shift Conflict"]))
    ][DISPLAY_COLUMNS].copy()

    if exceptions.empty:
        worksheet["A1"] = "No cycle or qualification exceptions found."
        worksheet["A1"].font = Font(bold=True)
        return

    write_dataframe(worksheet, exceptions)
    format_table_sheet(worksheet)


def create_roster_exceptions_sheet(workbook, roster_exceptions):
    worksheet = workbook.create_sheet("Roster Exceptions")

    if roster_exceptions.empty:
        worksheet["A1"] = "No Animation roster exceptions found."
        worksheet["A1"].font = Font(bold=True)
        return

    write_dataframe(worksheet, roster_exceptions)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="C65911")
        cell.font = Font(color="FFFFFF", bold=True)

    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 14
    worksheet.column_dimensions["D"].width = 31
    worksheet.column_dimensions["E"].width = 65

    for row in worksheet.iter_rows(min_row=2):
        row[2].number_format = "@"
        row[4].alignment = Alignment(wrap_text=True, vertical="top")


def create_run_summary(
    workbook,
    tracker_data,
    report_name,
    roster_count=None,
    matched_roster_count=None,
):
    worksheet = workbook.create_sheet("Run Summary")
    worksheet.sheet_view.showGridLines = False

    rows = [
        ("Report", report_name),
        ("Processing Date", TODAY.to_pydatetime()),
        (
            "Qualification Records",
            tracker_data[["SSO", "Qualification"]].drop_duplicates().shape[0],
        ),
        ("Annual Cycle Records", len(tracker_data)),
        ("Unique Technicians", tracker_data["SSO"].nunique()),
        (
            "Qualifications",
            ", ".join(sorted(tracker_data["Qualification"].unique())),
        ),
        (
            "Disciplines",
            ", ".join(sorted(tracker_data["Discipline"].unique())),
        ),
    ]

    if roster_count is not None:
        rows.append(("Unique Animation Roster SSOs", roster_count))
    if matched_roster_count is not None:
        rows.append(("Matched Official UCMOD4 SSOs", matched_roster_count))

    for row_number, (label, value) in enumerate(rows, start=1):
        worksheet.cell(row=row_number, column=1, value=label)
        worksheet.cell(row=row_number, column=2, value=value)
        worksheet.cell(row=row_number, column=1).font = Font(bold=True)

    worksheet["B2"].number_format = "m/d/yyyy"
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 60


def create_tracker_workbook(
    tracker_data,
    report_name,
    output_path,
    roster_exceptions=None,
    roster_count=None,
    matched_roster_count=None,
):
    workbook = Workbook()
    create_dashboard(workbook, tracker_data, report_name)

    status_sheet = workbook.create_sheet("PM Status")
    status_data = tracker_data[DISPLAY_COLUMNS].copy()
    write_dataframe(status_sheet, status_data)
    format_table_sheet(status_sheet)

    create_exceptions_sheet(workbook, tracker_data)

    if roster_exceptions is not None:
        create_roster_exceptions_sheet(workbook, roster_exceptions)

    create_run_summary(
        workbook,
        tracker_data,
        report_name,
        roster_count=roster_count,
        matched_roster_count=matched_roster_count,
    )

    workbook.save(output_path)


def main():
        started = time.perf_counter()

        print("=" * 70)
        print("PM Qualification Activity Tracker - Excel Report Generator")
        print("=" * 70)

        try:
            mapping = load_mapping()
            qualifications = read_qualifications(mapping)

            animation_ssos, roster_exceptions, roster_count = (
                read_animation_roster(qualifications)
            )

            tasks, qualifications_with_tasks = read_tasks()
            cycles = build_cycles(qualifications)

            results = assign_tasks_to_cycles(
                cycles,
                tasks,
                qualifications_with_tasks,
            )

            results, updated_history, history_stats = apply_cycle_history(
                results
            )

            results, schedule_path, schedule_stats = add_shift_to_results(
                results
            )

        except Exception as error:
            print(f"ERROR: {error}")
            return 1

        # Animation roster employees appear only in the Animation report.
        # The Animation report evaluates only UCMOD4.
        animation_data = results[
            results["SSO"].astype(str).isin(animation_ssos)
            & (results["Qualification"] == "UCMOD4")
        ].copy()

        animation_data["Venue"] = "Animation"

        # Remove every Animation roster SSO completely from the normal reports.
        standard_results = results[
            ~results["SSO"].astype(str).isin(animation_ssos)
        ].copy()

        OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

        standard_reports = {
            "Wolf": standard_results[
                standard_results["Venue"] == "Wolf"
            ].copy(),
            "Darkmoor": standard_results[
                standard_results["Venue"] == "Darkmoor"
            ].copy(),
            "Unchained": standard_results[
                standard_results["Venue"] == "Unchained"
            ].copy(),
        }

        try:
            for report_name, report_data in standard_reports.items():
                filename = OUTPUT_FILES[report_name]

                create_tracker_workbook(
                    report_data,
                    report_name,
                    OUTPUT_FOLDER / filename,
                )

                print(
                    f"Created: {filename} "
                    f"({len(report_data):,} cycle rows)"
                )

            animation_filename = OUTPUT_FILES["Animation"]

            create_tracker_workbook(
                animation_data,
                "Animation",
                OUTPUT_FOLDER / animation_filename,
                roster_exceptions=roster_exceptions,
                roster_count=roster_count,
                matched_roster_count=len(animation_ssos),
            )

            print(
                f"Created: {animation_filename} "
                f"({len(animation_data):,} cycle rows)"
            )

        except Exception as error:
            print()
            print(f"ERROR creating Excel reports: {error}")
            print("Cycle history was not updated.")
            return 1

        # Save historical evidence only after all four workbooks succeed.
        try:
            save_cycle_history(updated_history)
        except Exception as error:
            print()
            print(f"ERROR saving cycle history: {error}")
            print(
                "The Excel reports were created, but cycle history "
                "was not updated."
            )
            return 1

        print()
        print("Cycle history")
        print("-------------")

        for label, value in history_stats.items():
            print(f"{label}: {value:,}")

        print()
        print("Line Schedule")
        print("-------------")
        print(f"Source: {schedule_path}")
        for label, value in schedule_stats.items():
            print(f"{label}: {value:,}")

        print()
        print(f"Animation roster unique SSOs: {roster_count:,}")
        print(
            "Animation SSOs matched to UCMOD4: "
            f"{len(animation_ssos):,}"
        )
        print(
            f"Animation roster exceptions: "
            f"{len(roster_exceptions):,}"
        )
        print(f"Output folder: {OUTPUT_FOLDER}")
        print(
            "Total processing time: "
            f"{time.perf_counter() - started:,.1f} seconds"
        )
        print("Excel report generation completed successfully.")

        return 0


if __name__ == "__main__":
    sys.exit(main())