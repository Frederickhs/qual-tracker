from pathlib import Path
import warnings

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
TASK_FOLDER = PROJECT_ROOT / "Input" / "Tasks"
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}

# Maximo-generated workbooks may not contain Excel's normal default style.
# This warning does not affect the report data.
warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style*",
    category=UserWarning,
)


def get_task_files():
    return sorted(
        [
            file_path
            for file_path in TASK_FOLDER.iterdir()
            if file_path.is_file()
            and not file_path.name.startswith("~$")
            and file_path.suffix.lower() in SUPPORTED_EXTENSIONS
        ],
        key=lambda file_path: file_path.name.lower(),
    )


def normalize_header(value):
    if value is None:
        return ""
    return str(value).strip()


def find_header_row(worksheet, maximum_rows=15):
    """Find the Maximo header row instead of assuming it is row 1."""
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=maximum_rows, values_only=True),
        start=1,
    ):
        headers = {
            normalize_header(value): index
            for index, value in enumerate(row)
            if normalize_header(value)
        }

        if "Qualification" in headers and "Status" in headers:
            return row_number, headers

    return None, None


def inspect_workbook(file_path):
    print()
    print("=" * 70)
    print(f"File: {file_path.name}")
    print("=" * 70)

    workbook = load_workbook(file_path, read_only=True, data_only=True)

    qualifications = set()
    total_sheets = len(workbook.sheetnames)
    sheets_with_task_data = 0
    task_rows = 0
    completed_rows = 0
    first_header_row = None

    for worksheet in workbook.worksheets:
        worksheet.reset_dimensions()
        header_row_number, headers = find_header_row(worksheet)

        if headers is None:
            continue

        sheets_with_task_data += 1

        if first_header_row is None:
            first_header_row = header_row_number

        qualification_index = headers["Qualification"]
        status_index = headers["Status"]
        highest_needed_index = max(qualification_index, status_index)

        for row in worksheet.iter_rows(
            min_row=header_row_number + 1,
            values_only=True,
        ):
            if len(row) <= highest_needed_index:
                continue

            qualification = row[qualification_index]

            if qualification is None:
                continue

            qualification = str(qualification).strip().upper()

            if not qualification:
                continue

            qualifications.add(qualification)
            task_rows += 1

            status = row[status_index]

            if status is not None and str(status).strip().upper() == "COMP":
                completed_rows += 1

    workbook.close()

    print(f"Worksheets found: {total_sheets}")
    print(f"Worksheets with task data: {sheets_with_task_data}")
    print(f"First detected header row: {first_header_row or 'Not found'}")
    print(f"Task rows found: {task_rows:,}")
    print(f"COMP rows found: {completed_rows:,}")
    print("Qualifications found:")

    if qualifications:
        for qualification in sorted(qualifications):
            print(f"  - {qualification}")
    else:
        print("  None")


def main():
    task_files = get_task_files()

    print("=" * 70)
    print("PM Qualification Activity Tracker - Task Report Inspector")
    print("=" * 70)
    print(f"Task reports found: {len(task_files)}")

    if not task_files:
        print("No task reports are currently available.")
        return

    for file_path in task_files:
        inspect_workbook(file_path)

    print()
    print("Task report inspection completed successfully.")


if __name__ == "__main__":
    main()
